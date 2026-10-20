"""File-system tools used by the LLM-powered file assistant."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def _metadata(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime, tz=timezone.utc
        ).isoformat(),
        "extension": path.suffix.lower(),
    }


def _failure(filepath: str, error: str) -> dict[str, Any]:
    return {"success": False, "filepath": filepath, "error": error}


def _extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".txt":
        return path.read_text(encoding="utf-8")

    if extension == ".pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(str(path))
        except PdfReadError as exc:
            raise ValueError(f"Invalid PDF document: {exc}") from exc
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError

        try:
            document = Document(str(path))
        except PackageNotFoundError as exc:
            raise ValueError(f"Invalid DOCX document: {exc}") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        ]
        return "\n".join(part for part in [*paragraphs, *tables] if part)

    raise ValueError(
        f"Unsupported file type '{extension}'. Supported types: PDF, TXT, DOCX."
    )


def read_file(filepath: str) -> dict[str, Any]:
    """Extract text and metadata from a PDF, TXT, or DOCX file."""
    path = Path(filepath)
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return _failure(
            filepath,
            f"Unsupported file type '{path.suffix}'. Supported types: PDF, TXT, DOCX.",
        )
    if not path.is_file():
        return _failure(filepath, "File does not exist or is not a regular file.")

    try:
        content = _extract_text(path)
        return {
            "success": True,
            "filepath": str(path),
            "content": content,
            "metadata": _metadata(path),
        }
    except FileNotFoundError:
        return _failure(filepath, "File does not exist.")
    except PermissionError:
        return _failure(filepath, "Permission denied while reading the file.")
    except (OSError, UnicodeError, BadZipFile, ValueError) as exc:
        return _failure(filepath, f"Could not read file: {exc}")
    except ImportError as exc:
        return _failure(filepath, f"Missing file-reading dependency: {exc}")


def list_files(directory: str, extension: str | None = None) -> list[dict[str, Any]]:
    """List files directly inside a directory, optionally filtering by extension."""
    path = Path(directory)
    normalized_extension = None
    if extension:
        normalized_extension = extension.lower()
        if not normalized_extension.startswith("."):
            normalized_extension = f".{normalized_extension}"

    try:
        if not path.is_dir():
            return [{"success": False, "directory": directory, "error": "Directory does not exist."}]
        results = []
        for item in sorted(path.iterdir(), key=lambda entry: entry.name.lower()):
            if item.is_file() and (
                normalized_extension is None
                or item.suffix.lower() == normalized_extension
            ):
                results.append(_metadata(item))
        return results
    except (OSError, PermissionError) as exc:
        return [{"success": False, "directory": directory, "error": str(exc)}]


def write_file(filepath: str, content: str) -> dict[str, Any]:
    """Write UTF-8 text to a file, creating parent directories when necessary."""
    try:
        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return {
            "success": True,
            "filepath": str(path),
            "metadata": _metadata(path),
        }
    except (OSError, UnicodeError, TypeError) as exc:
        return _failure(filepath, f"Could not write file: {exc}")


def search_in_file(filepath: str, keyword: str) -> dict[str, Any]:
    """Find case-insensitive keyword occurrences and return their context."""
    if not keyword:
        return _failure(filepath, "Keyword must not be empty.")

    result = read_file(filepath)
    if not result["success"]:
        return result

    content = result["content"]
    folded_content = content.casefold()
    folded_keyword = keyword.casefold()
    matches = []
    position = folded_content.find(folded_keyword)
    while position != -1:
        start = max(0, position - 60)
        end = min(len(content), position + len(keyword) + 60)
        matches.append(
            {
                "position": position,
                "context": content[start:end],
            }
        )
        position = folded_content.find(folded_keyword, position + 1)

    return {
        "success": True,
        "filepath": result["filepath"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
        "metadata": result["metadata"],
    }
